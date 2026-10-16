"""Monta a pasta de um Space do Hugging Face com a demonstração do SGPP.

Uso (da raiz do projeto):
    python scripts/montar_demo_huggingface.py                # gera na pasta temporária (o caminho é mostrado)
    python scripts/montar_demo_huggingface.py --destino X    # gera em outra pasta
    python scripts/montar_demo_huggingface.py --alvo github --destino X   # repositório para o Streamlit Community Cloud

O que vai para a demonstração:
  * o código do app, os estilos e o brasão;
  * a planilha de exemplo do repositório com os nomes de gestores **anonimizados**;
  * contas de demonstração (uma por nível), com senha em hash e a senha de demonstração no README;
  * Dockerfile e README com o cabeçalho que o Hugging Face exige (o Streamlit roda via Docker).

O que NÃO vai: segredos (.streamlit/secrets.toml), a planilha real, dados internos e este script.
"""
import argparse
import os
import re
import shutil
import stat
import sys
import tempfile
import time
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
import contas  # noqa: E402

SENHA_DEMO = "demo-mpam-2026"
ARQUIVOS_DO_APP = ["app.py", "armazenamento.py", "contas.py", "dados_internos.py", "projetos.py", "requirements.txt"]
PLANILHA = "PROJETOS - Teste 1.xlsx"
PORTA = 7860
# Fora do OneDrive de propósito: a sincronização segura os arquivos e impede apagar a pasta numa nova geração
PASTA_PADRAO = Path(tempfile.gettempdir()) / "sgpp_demo_huggingface"

# (login, nome, nível, unidade). Os nomes "Membro N" são os que a planilha de exemplo já usa como gestores.
CONTAS_DEMO = [
    ("paula.plan", "Paula Planejamento Dias", "planejamento", "DPLAN"),
    ("carlos.comite", "Carlos Comitê Silva", "comite", "Gabinete da PGJ"),
    ("membro1", "Membro 1", "gestor_projeto", "NUPA"),
    ("membro2", "Membro 2", "gestor_unidade", "1ª PJ Parintins"),
    ("leitor.demo", "Leitor de Demonstração", "leitor", "DPLAN"),
]

_PADRAO_ANONIMO = re.compile(r"^Membros?\s+\w+$")


def anonimizar_gestores(caminho):
    """Troca por 'Membro N' todo gestor que não seja já um desses nomes genéricos. Devolve quantos foram trocados."""
    livro = openpyxl.load_workbook(caminho)
    folha = livro.active
    coluna = [c.value for c in folha[1]].index("Gestor") + 1
    alias, proximo, trocados = {}, 11, 0
    for linha in range(2, folha.max_row + 1):
        celula = folha.cell(row=linha, column=coluna)
        valor = " ".join(str(celula.value).split()) if celula.value else ""
        if not valor or _PADRAO_ANONIMO.match(valor):
            continue
        if valor not in alias:
            alias[valor] = f"Membro {proximo}"
            proximo += 1
        celula.value = alias[valor]
        trocados += 1
    livro.save(caminho)
    return trocados


def texto_usuarios_codigo(senha):
    entradas = []
    for login, nome, nivel, unidade in CONTAS_DEMO:
        texto, erro = contas.linha_usuario_codigo(login, nome, nivel, senha, unidade)
        if erro:
            raise ValueError(erro)
        entradas.append(texto)
    return '"""Usuários de DEMONSTRAÇÃO. Todos usam a mesma senha, que consta no README do Space."""\n\nUSUARIOS = {\n' + "\n".join(entradas) + "\n}\n"


DOCKERFILE = f"""FROM python:3.12-slim

# O Hugging Face roda o contêiner com o usuário de UID 1000
RUN useradd -m -u 1000 user
USER user
ENV HOME=/home/user \\
    PATH=/home/user/.local/bin:$PATH \\
    PYTHONUNBUFFERED=1 \\
    PIP_NO_CACHE_DIR=1
WORKDIR $HOME/app

COPY --chown=user requirements.txt .
RUN pip install --upgrade pip && pip install -r requirements.txt

COPY --chown=user . .

EXPOSE {PORTA}
CMD ["streamlit", "run", "app.py", "--server.port={PORTA}", "--server.address=0.0.0.0", "--server.headless=true", "--browser.gatherUsageStats=false"]
"""

CONFIG_SERVIDOR = """
# Só na demonstração: o Space abre o app dentro de um iframe, e a proteção contra XSRF/CORS do Streamlit atrapalha esse caso.
[server]
enableXsrfProtection = false
enableCORS = false
"""


def readme(senha, alvo="huggingface"):
    linhas_contas = "\n".join(
        f"| `{login}` | {contas.ROTULOS_NIVEL[nivel]} | {nome} |" for login, nome, nivel, _ in CONTAS_DEMO)
    cabecalho = f"""---
title: SGPP MPAM (demonstração)
colorFrom: red
colorTo: yellow
sdk: docker
app_port: {PORTA}
short_description: Demonstração do acompanhamento de projetos da DPLAN/MPAM
pinned: false
---

""" if alvo == "huggingface" else ""
    onde = "o Space reinicia" if alvo == "huggingface" else "o app reinicia ou é reimplantado"
    return cabecalho + f"""# SGPP MPAM: demonstração

Demonstração do Sistema de Gestão de Projetos da Diretoria de Planejamento (DPLAN) do Ministério Público do Estado do Amazonas.

**Os dados são de exemplo.** A planilha é uma amostra pequena com os nomes dos gestores anonimizados. Tudo o que for cadastrado, comentado ou alterado aqui **é apagado quando {onde}**.

## Como entrar

A senha de todas as contas de demonstração é `{senha}`.

| Login | Nível | Nome |
|---|---|---|
{linhas_contas}

## Roteiro sugerido

1. Entre como `membro2`, abra **Elaboração**, preencha um projeto e clique em **Enviar para a Planejamento**.
2. Saia e entre como `paula.plan`. Em **Revisão**, abra **Ver o projeto inteiro**, escreva uma orientação e clique em **Devolver para ajustes**.
3. Volte como `membro2`. Em **Meus Projetos**, leia o retorno da Planejamento e responda.
4. Como `paula.plan`, veja o **Quadro** (mova um cartão), as **Pendências** e a aba **Com o gestor** da Revisão.
5. Em **Acompanhamento**, registre uma atualização e veja a linha do tempo do projeto.
""" + ("" if alvo == "huggingface" else PUBLICACAO_GITHUB)


PUBLICACAO_GITHUB = """
## Como colocar no ar (Streamlit Community Cloud)

1. Em share.streamlit.io, entre com a conta do GitHub e clique em **Create app**.
2. Escolha este repositório, o branch `main` e o arquivo principal `app.py`.
3. Em **Advanced settings**, escolha Python 3.12. Não é preciso configurar Secrets: as contas de demonstração já estão no código.
4. O app dorme depois de cerca de 12 horas sem acesso; abra o link alguns minutos antes de usar.

Este repositório não contém dado real, segredo nem credencial do Entra ID. Para usar o app de verdade, com a planilha do SharePoint, veja o repositório principal.
"""


def _apagar(pasta, tentativas=5):
    """Apaga a pasta com novas tentativas (no Windows, antivírus e sincronização seguram arquivos por instantes)."""
    def liberar_e_repetir(funcao, caminho, _):
        os.chmod(caminho, stat.S_IWRITE)
        funcao(caminho)

    for numero in range(tentativas):
        try:
            shutil.rmtree(pasta, onerror=liberar_e_repetir)
            return
        except OSError:
            if numero == tentativas - 1:
                raise OSError(f"Não consegui apagar {pasta}. Feche o que estiver usando essa pasta e tente de novo.") from None
            time.sleep(0.6)


def montar(destino, senha=SENHA_DEMO, alvo="huggingface"):
    """Gera a pasta da demonstração em `destino` (apagada antes, se existir). `alvo` é "huggingface" (Docker) ou "github"
    (repositório para o Streamlit Community Cloud). Devolve o caminho."""
    destino = Path(destino)
    if destino.exists():
        _apagar(destino)
    destino.mkdir(parents=True)
    for nome in ARQUIVOS_DO_APP:
        shutil.copy(RAIZ / nome, destino / nome)
    shutil.copytree(RAIZ / "assets", destino / "assets")

    shutil.copy(RAIZ / PLANILHA, destino / PLANILHA)
    anonimizar_gestores(destino / PLANILHA)

    (destino / "usuarios_codigo.py").write_text(texto_usuarios_codigo(senha), encoding="utf-8")
    pasta_streamlit = destino / ".streamlit"
    pasta_streamlit.mkdir()
    configuracao = (RAIZ / ".streamlit" / "config.toml").read_text(encoding="utf-8")
    if alvo == "huggingface":
        configuracao += CONFIG_SERVIDOR
        (destino / "Dockerfile").write_text(DOCKERFILE, encoding="utf-8", newline="\n")
    (pasta_streamlit / "config.toml").write_text(configuracao, encoding="utf-8")
    (destino / "README.md").write_text(readme(senha, alvo), encoding="utf-8")
    ignorados = "__pycache__/\n*.pyc\n.streamlit/secrets.toml\nDADOS_INTERNOS - SGPP.xlsx\n"
    (destino / (".dockerignore" if alvo == "huggingface" else ".gitignore")).write_text(ignorados, encoding="utf-8")
    return destino


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--destino", default=str(PASTA_PADRAO), help="pasta de saída")
    parser.add_argument("--alvo", choices=["huggingface", "github"], default="huggingface",
                        help="huggingface (Space Docker) ou github (repositório para o Streamlit Community Cloud)")
    argumentos = parser.parse_args()
    pasta = montar(argumentos.destino, alvo=argumentos.alvo)
    print(f"Pasta da demonstração gerada em: {pasta}")
    print(f"Senha das contas de demonstração: {SENHA_DEMO}")


if __name__ == "__main__":
    main()
