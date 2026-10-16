import os
import re
import time

import openpyxl
import pandas as pd

import contas
import dados_internos
from armazenamento import ArmazenamentoLocal
from apoio import (
    CAMINHO_INTERNO, SENHA, abrir_app, botao, campo, entrar, preencher_rascunho, semear_usuarios, textos_visiveis,
)
from conftest import ADMIN, PLANILHA_EXEMPLO

MENU = {
    "dashboard": ":material/dashboard: Dashboard",
    "meus": ":material/folder_open: Meus Projetos",
    "elaboracao": ":material/edit_note: Elaboração",
    "acompanhamento": ":material/timeline: Acompanhamento",
    "planilha": ":material/table_view: Planilha Central",
    "exportar": ":material/download: Exportar",
    "revisao": ":material/rate_review: Revisão",
    "painel": ":material/admin_panel_settings: Painel de Gestão",
    "usuarios": ":material/group: Usuários e Níveis",
}


def ir_para(at, pagina):
    at.sidebar.radio[0].set_value(MENU[pagina])
    return at.run()


def entrar_como_admin(pasta):
    at = abrir_app()
    return entrar(at, *ADMIN)


def internos():
    abas, _ = dados_internos.ler_internos(ArmazenamentoLocal(CAMINHO_INTERNO))
    return abas


# ------------------------------------------------------------------ login

def test_login_nao_expoe_senhas_e_so_entra_com_credenciais_validas(pasta):
    at = abrir_app()
    tela = textos_visiveis(at)
    for antiga in ("dplan2026", "mpam2026", "123"):
        assert antiga not in tela
    assert {t.label for t in at.text_input} >= {"Login institucional", "Senha"}

    at = entrar(at, "chefe.geral", "senha-errada")
    assert not at.session_state["autenticado"]
    assert any("incorretos" in e.value for e in at.error)

    for login, senha in [("admin", "123"), ("membro1", "mpam2026"), ("qualquer.um", "dplan2026")]:
        assert not entrar(abrir_app(), login, senha).session_state["autenticado"]


def test_administrador_inicial_entra_como_gestor_geral_sem_conta_na_planilha(pasta):
    at = entrar_como_admin(pasta)
    assert at.session_state["autenticado"]
    assert at.session_state["nivel"] == "administrador" and at.session_state["perfil"] == "Gestor_Geral"
    assert not at.session_state["conta_na_planilha"]
    assert not at.exception
    assert MENU["usuarios"] in at.sidebar.radio[0].options and MENU["painel"] in at.sidebar.radio[0].options


def test_cinco_senhas_erradas_bloqueiam_o_login_ate_para_a_senha_certa(pasta):
    semear_usuarios(("ana.souza", "Ana Souza Lima", "leitor"))
    at = abrir_app()
    for _ in range(5):
        at = entrar(at, "ana.souza", "errada-errada-1")
    at = entrar(at, "ana.souza", SENHA)
    assert not at.session_state["autenticado"]
    assert any("Muitas tentativas" in e.value for e in at.error)
    # outro login não é afetado
    assert entrar(abrir_app(), *ADMIN).session_state["autenticado"]


# ------------------------------------------------ cadastro, aprovação e níveis

def solicitar(at, login="bia.santos", nome="Beatriz Santos Melo", senha=SENHA):
    campo(at, "Login institucional (ex.: nome.sobrenome)").set_value(login)
    campo(at, "Nome completo (igual ao que consta como Gestor nos projetos)").set_value(nome)
    campo(at, f"E-mail institucional (@{contas.DOMINIO_EMAIL})").set_value(f"{login}@mpam.mp.br")
    campo(at, "Unidade / Promotoria").set_value("1ª PJ Parintins")
    campo(at, f"Senha (mínimo de {contas.TAMANHO_MINIMO_SENHA} caracteres)").set_value(senha)
    campo(at, "Repita a senha").set_value(senha)
    return botao(at, "Solicitar cadastro").click().run()


def test_fluxo_completo_cadastro_aprovacao_e_primeiro_acesso(pasta):
    at = solicitar(abrir_app())
    assert any("Cadastro enviado" in s.value for s in at.success)
    conta = contas.buscar_conta(internos()["Usuarios"], "bia.santos")
    assert conta["situacao"] == "pendente" and conta["senha_hash"].startswith("$argon2id$")

    # pendente não entra
    at = entrar(abrir_app(), "bia.santos", SENHA)
    assert not at.session_state["autenticado"]
    assert any("aguarda" in e.value for e in at.error)

    # a Planejamento aprova com um nível
    at = ir_para(entrar_como_admin(pasta), "usuarios")
    assert not at.exception
    assert any("Aguardando aprovação (1)" in t.label for t in at.tabs)
    at.selectbox(key="nivel_pend_bia.santos").set_value("Gestor de unidade")
    at = botao(at, "Aprovar").click().run()
    assert not at.exception
    assert any("aprovado" in s.value for s in at.success)
    conta = contas.buscar_conta(internos()["Usuarios"], "bia.santos")
    assert conta["situacao"] == "ativo" and conta["nivel"] == "gestor_unidade" and conta["decidido_por"] == ADMIN[0]
    assert any(r.acao == "cadastro_aprovado" and r.alvo == "bia.santos" for r in internos()["Registro"].itertuples())

    # agora entra, sem acesso às telas de gestão
    at = entrar(abrir_app(), "bia.santos", SENHA)
    assert at.session_state["autenticado"] and at.session_state["nivel"] == "gestor_unidade"
    opcoes = at.sidebar.radio[0].options
    assert MENU["usuarios"] not in opcoes and MENU["painel"] not in opcoes
    assert not at.exception


def test_comite_so_concede_niveis_abaixo_do_proprio(pasta):
    semear_usuarios(("carlos.comite", "Carlos Comitê Silva", "comite"))
    at = solicitar(abrir_app())
    at = entrar(abrir_app(), "carlos.comite", SENHA)
    assert at.session_state["perfil"] == "Comite"
    assert MENU["usuarios"] in at.sidebar.radio[0].options and MENU["painel"] not in at.sidebar.radio[0].options
    at = ir_para(at, "usuarios")
    niveis = at.selectbox(key="nivel_pend_bia.santos").options
    assert "Leitor" in niveis and "Gestor de unidade" in niveis
    assert "Planejamento" not in niveis and "Comitê" not in niveis and "Administrador" not in niveis


def test_planejamento_nao_concede_administrador(pasta):
    semear_usuarios(("paula.plan", "Paula Planejamento Dias", "planejamento"))
    solicitar(abrir_app())
    at = ir_para(entrar(abrir_app(), "paula.plan", SENHA), "usuarios")
    niveis = at.selectbox(key="nivel_pend_bia.santos").options
    assert "Comitê" in niveis and "Planejamento" in niveis and "Administrador" not in niveis


def test_senha_temporaria_obriga_a_troca_antes_de_qualquer_tela(pasta):
    semear_usuarios(("ana.souza", "Ana Souza Lima", "gestor_projeto", True))
    at = entrar(abrir_app(), "ana.souza", SENHA)
    assert at.session_state["autenticado"] and at.session_state["trocar_senha"]
    assert not at.sidebar.radio  # sem menu enquanto não trocar
    campo(at, "Senha temporária recebida").set_value(SENHA)
    campo(at, f"Nova senha (mínimo de {contas.TAMANHO_MINIMO_SENHA} caracteres)").set_value("senha-nova-bem-longa-9")
    campo(at, "Repita a nova senha").set_value("senha-nova-bem-longa-9")
    at = botao(at, "Alterar senha").click().run()
    assert not at.exception
    assert not at.session_state["trocar_senha"]
    assert at.sidebar.radio  # menu liberado
    assert entrar(abrir_app(), "ana.souza", "senha-nova-bem-longa-9").session_state["autenticado"]
    assert not entrar(abrir_app(), "ana.souza", SENHA).session_state["autenticado"]


def test_desativar_conta_derruba_o_acesso(pasta):
    semear_usuarios(("ana.souza", "Ana Souza Lima", "leitor"))
    at = ir_para(entrar_como_admin(pasta), "usuarios")
    at.selectbox(key="conta_alvo").set_value("ana.souza")
    at = at.run()
    at.radio(key="acao_ana.souza").set_value("Desativar")
    at = at.run()
    at = botao(at, "Desativar conta").click().run()
    assert not at.exception
    assert contas.buscar_conta(internos()["Usuarios"], "ana.souza")["situacao"] == "desativado"
    at = entrar(abrir_app(), "ana.souza", SENHA)
    assert not at.session_state["autenticado"]
    assert any("desativada" in e.value for e in at.error)


# ------------------------------------------------------------------ telas

def test_nenhuma_tela_menciona_auditoria(pasta):
    at = entrar_como_admin(pasta)
    for pagina in MENU:
        at = ir_para(at, pagina)
        assert not at.exception, pagina
        assert not re.search(r"audit", textos_visiveis(at), re.IGNORECASE), pagina


def test_filtro_de_institucionalizados_reduz_a_tabela(pasta):
    # No exemplo todos são "Sim"; marca 5 como "Não" na cópia temporária para o filtro ter o que separar
    livro = openpyxl.load_workbook(PLANILHA_EXEMPLO)
    folha = livro.active
    coluna = [c.value for c in folha[1]].index("Institucionalizado") + 1
    for linha in range(2, 7):
        folha.cell(row=linha, column=coluna, value="Não")
    livro.save(PLANILHA_EXEMPLO)

    base = pd.read_excel(PLANILHA_EXEMPLO)
    # O painel só lista projetos com Área, Status e Município preenchidos (filtros do próprio painel)
    base = base.dropna(subset=["Área de Atuação", "Status", "Município"])
    total = len(base)
    sim = int((base["Institucionalizado"].astype(str).str.strip().str.lower() == "sim").sum())
    assert 0 < sim < total

    at = ir_para(entrar_como_admin(pasta), "dashboard")
    at.checkbox(key="inst_dash").set_value(True)
    at = at.run()
    assert not at.exception
    assert any(f"Mostrando {sim} de {total}" in c.value for c in at.caption)


def test_filtro_por_coluna_e_busca_livre(pasta):
    at = ir_para(entrar_como_admin(pasta), "dashboard")
    area = at.multiselect(key="filtro_dash_Área de Atuação")
    escolhida = area.options[0]
    area.set_value([escolhida])
    at = at.run()
    assert any("Mostrando" in c.value for c in at.caption)
    at.multiselect(key="filtro_dash_Área de Atuação").set_value([])
    at.text_input(key="busca_dash").set_value("texto-que-nao-existe-em-lugar-nenhum")
    at = at.run()
    assert any("Mostrando 0 de" in c.value for c in at.caption)


# ----------------------------------------- fluxos de projetos (regressão)

def cabecalho(caminho):
    return [c.value for c in openpyxl.load_workbook(caminho).active[1]]


def test_envio_exclusao_e_conflito_de_projetos_continuam_funcionando(pasta):
    original = cabecalho(PLANILHA_EXEMPLO)
    linhas = openpyxl.load_workbook(PLANILHA_EXEMPLO).active.max_row

    at = ir_para(entrar_como_admin(pasta), "elaboracao")
    at = botao(preencher_rascunho(at, nome="Projeto de Teste Automatizado", unidade="DPLAN", ementa="Ementa de teste"),
               "Enviar para a Planejamento").click().run()
    assert not at.exception
    assert any("Projeto enviado" in s.value for s in at.success)
    planilha = openpyxl.load_workbook(PLANILHA_EXEMPLO).active
    assert planilha.max_row == linhas + 1
    assert [c.value for c in planilha[1]][: len(original)] == original  # cabeçalho que o Power BI lê

    at = ir_para(at, "painel")
    assert not at.exception
    selecao = at.selectbox(key="sel_del_admin")
    novo_id = pd.read_excel(PLANILHA_EXEMPLO, dtype=str).set_index("Iniciativa.1").loc["Projeto de Teste Automatizado", "ID_PROJETO"]
    selecao.set_value(novo_id)
    at.text_input(key="mot_del_admin").set_value("Teste automatizado de exclusão")
    at = botao(at, "Executar Exclusão").click().run()
    assert not at.exception
    cab = cabecalho(PLANILHA_EXEMPLO)
    achou = [r for r in openpyxl.load_workbook(PLANILHA_EXEMPLO).active.iter_rows(min_row=2, values_only=True)
             if r[cab.index("Iniciativa", 1)] == "Projeto de Teste Automatizado"]
    assert achou and achou[0][cab.index("SITUACAO_REGISTRO")] == "Excluído"
    # a planilha muda por fora entre a tela e o clique em salvar
    at = ir_para(at, "painel")
    time.sleep(0.05)
    os.utime(PLANILHA_EXEMPLO)
    at = botao(at, "Gravar Avaliações").click().run()
    assert any("alterada por outra pessoa" in e.value for e in at.error)
