"""O pacote do Space do Hugging Face: o que contém, o que não pode conter e se o app sobe e funciona sozinho nele."""
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import montar_demo_huggingface as demo  # noqa: E402

NOMES_REAIS = ("Yara", "Romina", "Vitor Moreira", "Albuquerque", "Brito Carvalho")


@pytest.fixture(scope="module")
def pacote(tmp_path_factory):
    return demo.montar(tmp_path_factory.mktemp("space") / "huggingface")


def texto_de(pacote, nome):
    return (pacote / nome).read_text(encoding="utf-8")


def test_readme_tem_o_cabecalho_de_space_docker_e_a_senha_de_demonstracao(pacote):
    readme = texto_de(pacote, "README.md")
    cabecalho = re.match(r"^---\n(.*?)\n---\n", readme, re.DOTALL).group(1)
    campos = dict(linha.split(": ", 1) for linha in cabecalho.splitlines())
    assert campos["sdk"] == "docker" and campos["app_port"] == "7860" and "SGPP" in campos["title"]
    assert "emoji" not in campos  # o projeto não usa emoji
    assert demo.SENHA_DEMO in readme
    for login, *_ in demo.CONTAS_DEMO:
        assert f"`{login}`" in readme


def test_dockerfile_segue_a_receita_do_hugging_face(pacote):
    docker = texto_de(pacote, "Dockerfile")
    assert "useradd -m -u 1000 user" in docker and "USER user" in docker
    assert "COPY --chown=user" in docker and "python:3.12" in docker
    assert "--server.port=7860" in docker and "--server.address=0.0.0.0" in docker and "EXPOSE 7860" in docker
    assert "\r" not in docker  # fim de linha do Linux


def test_pacote_nao_leva_segredos_nem_dados_internos_nem_o_proprio_empacotador(pacote):
    nomes = {p.relative_to(pacote).as_posix() for p in pacote.rglob("*") if p.is_file()}
    assert ".streamlit/secrets.toml" not in nomes and ".streamlit/config.toml" in nomes
    assert not [n for n in nomes if "DADOS_INTERNOS" in n or n.startswith(("tests/", "scripts/", "build/"))]
    todo_texto = "\n".join(texto_de(pacote, n) for n in nomes if n.endswith((".py", ".md", ".toml", ".txt", "Dockerfile")))
    assert "5lw8Q" not in todo_texto and "client_secret = \"" not in todo_texto.replace('client_secret = ""', "")


def test_usuarios_da_demonstracao_tem_hash_e_nunca_a_senha_em_texto(pacote):
    usuarios = texto_de(pacote, "usuarios_codigo.py")
    assert usuarios.count("$argon2id$") == len(demo.CONTAS_DEMO)
    assert demo.SENHA_DEMO not in usuarios


def test_planilha_da_demonstracao_nao_tem_nomes_reais(pacote):
    folha = openpyxl.load_workbook(pacote / demo.PLANILHA).active
    cabecalho = [c.value for c in folha[1]]
    gestores = {" ".join(str(l[cabecalho.index("Gestor")]).split()) for l in folha.iter_rows(min_row=2, values_only=True) if l[cabecalho.index("Gestor")]}
    assert gestores and all(re.match(r"^Membros?\s+\w+$", g) for g in gestores), gestores
    todas = " ".join(str(c) for linha in folha.iter_rows(values_only=True) for c in linha if c)
    assert not [n for n in NOMES_REAIS if n in todas]


def test_anonimizar_gestores_troca_so_nomes_proprios_e_mantem_o_mesmo_apelido(tmp_path):
    caminho = tmp_path / "t.xlsx"
    livro = openpyxl.Workbook()
    livro.active.append(["Nº", "Gestor"])
    for i, g in enumerate(["Membro 1", "Maria da Silva Souza", None, "Maria da Silva  Souza", "João Pereira", "Membros I"], start=1):
        livro.active.append([i, g])
    livro.save(caminho)
    assert demo.anonimizar_gestores(caminho) == 3
    valores = [r[1] for r in openpyxl.load_workbook(caminho).active.iter_rows(min_row=2, values_only=True)]
    assert valores == ["Membro 1", "Membro 11", None, "Membro 11", "Membro 12", "Membros I"]


# ------------------------------------------------------------------ publicação (sem rede: API simulada)

class ApiFalsa:
    def __init__(self):
        self.chamadas = []

    def whoami(self):
        self.chamadas.append(("whoami",))
        return {"name": "Hugo.Rocha"}

    def create_repo(self, **kw):
        self.chamadas.append(("create_repo", kw))

    def upload_folder(self, **kw):
        self.chamadas.append(("upload_folder", kw))


def test_simular_nao_chama_a_api_nem_envia_nada(tmp_path):
    import publicar_demo_huggingface as publicar

    api, falas = ApiFalsa(), []
    assert publicar.publicar(simular=True, api=api, pasta=tmp_path / "p", saida=falas.append) is None
    assert api.chamadas == [] and any("nada foi enviado" in f for f in falas)
    assert (tmp_path / "p" / "Dockerfile").exists()  # o pacote é gerado mesmo na simulação


def test_publicar_cria_space_docker_privado_por_padrao_e_envia_a_pasta(tmp_path):
    import publicar_demo_huggingface as publicar

    api, falas = ApiFalsa(), []
    repo = publicar.publicar(api=api, pasta=tmp_path / "p", saida=falas.append)
    assert repo == "Hugo.Rocha/sgpp-mpam-demo"
    criar = next(c for c in api.chamadas if c[0] == "create_repo")[1]
    assert criar == {"repo_id": repo, "repo_type": "space", "space_sdk": "docker", "private": True, "exist_ok": True}
    enviar = next(c for c in api.chamadas if c[0] == "upload_folder")[1]
    assert enviar["folder_path"] == str(tmp_path / "p") and enviar["repo_id"] == repo and enviar["repo_type"] == "space"
    assert any("hugo-rocha-sgpp-mpam-demo.hf.space" in f for f in falas)  # endereço direto do app
    assert any("privada" in f for f in falas)


def test_publico_so_com_a_opcao_explicita(tmp_path):
    import publicar_demo_huggingface as publicar

    api = ApiFalsa()
    publicar.publicar(publico=True, api=api, pasta=tmp_path / "p", saida=lambda _: None)
    assert next(c for c in api.chamadas if c[0] == "create_repo")[1]["private"] is False


DRIVER = """
import json, os, sys
from streamlit.testing.v1 import AppTest

senha, logins = sys.argv[1], sys.argv[2:]
resultado = {}
for login in logins:
    at = AppTest.from_file(os.path.join(os.getcwd(), "app.py"), default_timeout=120).run()
    campos = {t.label: t for t in at.text_input}
    campos["Login institucional"].set_value(login)
    campos["Senha"].set_value(senha)
    at = next(b for b in at.button if "Entrar" in b.label).click().run()
    entrou = bool(at.session_state["autenticado"])
    excecoes, paginas = [], 0
    if entrou:
        for opcao in list(at.sidebar.radio[0].options):
            at.sidebar.radio[0].set_value(opcao)
            at = at.run()
            paginas += 1
            excecoes += [e.value for e in at.exception]
    resultado[login] = {"entrou": entrou, "nivel": at.session_state["nivel"] if entrou else None, "paginas": paginas, "excecoes": excecoes}
print("RESULTADO=" + json.dumps(resultado))
"""


def test_app_do_pacote_sobe_sozinho_sem_segredos_e_todos_os_perfis_entram(pacote, tmp_path):
    driver = tmp_path / "driver.py"
    driver.write_text(DRIVER, encoding="utf-8")
    logins = [c[0] for c in demo.CONTAS_DEMO]
    saida = subprocess.run([sys.executable, str(driver), demo.SENHA_DEMO, *logins], cwd=pacote, capture_output=True, text=True, timeout=600)
    linha = next((l for l in saida.stdout.splitlines() if l.startswith("RESULTADO=")), None)
    assert linha, saida.stderr[-2000:]
    resultado = json.loads(linha.removeprefix("RESULTADO="))
    for login, _, nivel, _ in demo.CONTAS_DEMO:
        r = resultado[login]
        assert r["entrou"] and r["nivel"] == nivel, (login, r)
        assert r["paginas"] >= 4 and not r["excecoes"], (login, r)


# ------------------------------------------------------------------ alvo "github" (Streamlit Community Cloud)

def test_pacote_github_nao_tem_cabecalho_do_hugging_face_nem_dockerfile_e_nao_ignora_o_que_o_app_precisa(tmp_path):
    pasta = demo.montar(tmp_path / "gh", alvo="github")
    nomes = {p.relative_to(pasta).as_posix() for p in pasta.rglob("*") if p.is_file()}
    assert "Dockerfile" not in nomes and ".dockerignore" not in nomes and ".gitignore" in nomes
    readme = (pasta / "README.md").read_text(encoding="utf-8")
    assert not readme.startswith("---") and "Streamlit Community Cloud" in readme and demo.SENHA_DEMO in readme
    assert "enableXsrfProtection" not in (pasta / ".streamlit" / "config.toml").read_text(encoding="utf-8")
    ignorados = [l for l in (pasta / ".gitignore").read_text(encoding="utf-8").splitlines() if l]
    assert ".streamlit/secrets.toml" in ignorados
    # nada do que o app precisa para subir pode estar na lista de ignorados
    assert not [n for n in ("app.py", "usuarios_codigo.py", demo.PLANILHA, "requirements.txt", ".streamlit/config.toml") if n in ignorados]
