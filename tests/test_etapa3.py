"""Quadro (kanban) e relatório de pendências, pela interface."""
import re

import openpyxl
import pandas as pd
import pytest

import projetos
from apoio import entrar_codigo, semear_codigo, textos_visiveis
from conftest import PLANILHA_EXEMPLO
from test_app import MENU, ir_para
from test_etapa2 import ANA, CARLOS, LEITORA, PAULA, enviar_projeto, internos, projeto_novo

MENU.setdefault("quadro", ":material/view_kanban: Quadro")
MENU.setdefault("pendencias", ":material/checklist: Pendências")
STATUS_INDICE = {s: i for i, s in enumerate(projetos.STATUS)}


@pytest.fixture
def equipe(pasta, monkeypatch):
    semear_codigo(monkeypatch, ANA, PAULA, CARLOS, LEITORA, ("membro.dois", "Membro 2", "gestor_unidade"))


def amostra():
    return pd.read_excel(PLANILHA_EXEMPLO).rename(columns={"Iniciativa.1": "Nome da Iniciativa"})


def alterar_planilha(por_linha):
    """Edita a planilha de exemplo (cópia temporária): {n: {coluna: valor}} usando o número do projeto."""
    livro = openpyxl.load_workbook(PLANILHA_EXEMPLO)
    folha = livro.active
    cab = [c.value for c in folha[1]]
    n_col = cab.index("Nº") + 1
    for linha in range(2, folha.max_row + 1):
        n = folha.cell(row=linha, column=n_col).value
        for coluna, valor in por_linha.get(n, {}).items():
            folha.cell(row=linha, column=cab.index(coluna) + 1).value = valor  # cell(value=None) não limpa a célula
    livro.save(PLANILHA_EXEMPLO)


def colunas_do_quadro(at):
    """{coluna: quantidade} lido dos títulos do quadro."""
    html = "\n".join(str(m.value) for m in at.markdown)
    return {nome.strip(): int(qtd) for nome, qtd in re.findall(r"<div class='quadro-titulo'>(.*?)<span>(\d+)</span>", html)}


# ------------------------------------------------------------------ quadro

def test_quadro_aparece_para_todos_e_pendencias_so_para_coordenacao(equipe):
    for conta, tem_pendencias in ((ANA, False), (PAULA, True), (CARLOS, True), (LEITORA, False)):
        opcoes = entrar_codigo(conta[0]).sidebar.radio[0].options
        assert MENU["quadro"] in opcoes, conta
        assert (MENU["pendencias"] in opcoes) is tem_pendencias, conta


def test_quadro_tem_cinco_colunas_e_conta_todos_os_projetos_ativos(equipe):
    alterar_planilha({1: {"Status": "Há intenção formalizada"}, 2: {"Status": "Encerrado"}, 3: {"Status": None}})
    esperado = {}
    for status in amostra()["Status"]:
        coluna = projetos.coluna_do_quadro(status)
        esperado[coluna] = esperado.get(coluna, 0) + 1

    at = ir_para(entrar_codigo(PAULA[0]), "quadro")
    assert not at.exception
    contagens = colunas_do_quadro(at)
    assert list(contagens) == projetos.COLUNAS_QUADRO
    assert contagens == {c: esperado.get(c, 0) for c in projetos.COLUNAS_QUADRO}
    assert contagens["Outros status"] >= 3 and sum(contagens.values()) == len(amostra())


def test_cartao_de_vocabulario_antigo_mostra_o_status_original(equipe):
    alterar_planilha({1: {"Status": "Há intenção formalizada"}})
    at = ir_para(entrar_codigo(PAULA[0]), "quadro")
    html = "\n".join(str(m.value) for m in at.markdown)
    assert "Há intenção formalizada" in html


def test_planejamento_move_o_cartao_e_o_historico_registra(equipe):
    base = amostra()
    n = str(int(base[base["Status"] == "Em execução"].iloc[0]["Nº"]))
    at = ir_para(entrar_codigo(PAULA[0]), "quadro")
    antes = colunas_do_quadro(at)
    at = at.button(key=f"mv_{n}_{STATUS_INDICE['Concluído']}").click().run()
    assert not at.exception and any("movido para Concluído" in s.value for s in at.success)

    novo = pd.read_excel(PLANILHA_EXEMPLO)
    assert novo[novo["Nº"] == int(n)].iloc[0]["Status"] == "Concluído"
    evento = internos()["Historico"].iloc[-1]
    assert evento["evento"] == "editado" and evento["projeto_n"] == n
    assert evento["detalhe"] == "Status: de Em execução para Concluído" and evento["ator_login"] == "paula.plan"

    depois = colunas_do_quadro(at)
    assert depois["Em execução"] == antes["Em execução"] - 1 and depois["Concluído"] == antes["Concluído"] + 1


def test_cartao_nao_oferece_o_status_em_que_ja_esta(equipe):
    base = amostra()
    n = str(int(base[base["Status"] == "Em execução"].iloc[0]["Nº"]))
    at = ir_para(entrar_codigo(PAULA[0]), "quadro")
    chaves = {b.key for b in at.button if b.key and b.key.startswith(f"mv_{n}_")}
    assert chaves == {f"mv_{n}_{i}" for i, s in enumerate(projetos.STATUS) if s != "Em execução"}


def test_gestor_so_move_os_proprios_cartoes(equipe):
    base = amostra()
    meus = {str(int(n)) for n in base[base["Gestor"] == "Membro 2"]["Nº"]}
    de_outros = {str(int(n)) for n in base[base["Gestor"] != "Membro 2"]["Nº"]}
    assert meus and de_outros
    at = ir_para(entrar_codigo("membro.dois"), "quadro")
    chaves = {b.key for b in at.button if b.key and b.key.startswith("mv_")}
    assert {k.split("_")[1] for k in chaves} == meus  # só os dele

    n = sorted(meus)[0]
    destino = next(s for s in projetos.STATUS if s != projetos.texto(base[base["Nº"] == int(n)].iloc[0]["Status"]))
    at = at.button(key=f"mv_{n}_{STATUS_INDICE[destino]}").click().run()
    assert not at.exception and pd.read_excel(PLANILHA_EXEMPLO).set_index("Nº").loc[int(n), "Status"] == destino


def test_comite_e_leitor_consultam_mas_nao_movem(equipe):
    for conta in (CARLOS, LEITORA):
        at = ir_para(entrar_codigo(conta[0]), "quadro")
        assert not at.exception and not [b for b in at.button if b.key and b.key.startswith("mv_")]
        assert any("não tem projetos que possa mover" in c.value for c in at.caption)


def test_projeto_travado_so_se_move_pela_planejamento_ou_depois_de_liberado(equipe):
    enviar_projeto()
    n = str(int(float(projeto_novo()["Nº"])))
    # a gestora não vê o menu de mover num projeto travado
    at = ir_para(entrar_codigo(ANA[0]), "quadro")
    assert not [b for b in at.button if b.key and b.key.startswith(f"mv_{n}_")]
    # a Planejamento move sem alterar a trava
    at = ir_para(entrar_codigo(PAULA[0]), "quadro")
    at = at.button(key=f"mv_{n}_{STATUS_INDICE['Em execução']}").click().run()
    assert projeto_novo()["Status"] == "Em execução" and projeto_novo()["EDICAO"] == "Travada"
    # liberada, a gestora move e o projeto volta a ficar travado
    ir_para(entrar_codigo(PAULA[0]), "revisao")
    at = ir_para(entrar_codigo(PAULA[0]), "revisao")
    at.selectbox(key="revisao_alvo").set_value(n)
    at = at.run()
    at.button(key="liberar_direto").click().run()
    at = ir_para(entrar_codigo(ANA[0]), "quadro")
    at = at.button(key=f"mv_{n}_{STATUS_INDICE['Concluído']}").click().run()
    novo = projeto_novo()
    assert novo["Status"] == "Concluído" and novo["EDICAO"] == "Travada"
    assert [e for e in internos()["Historico"]["evento"]][-2:] == ["editado", "travado"]


# ------------------------------------------------------------------ pendências

def test_pendencias_contam_os_campos_obrigatorios_em_branco(equipe):
    alterar_planilha({1: {"Município": None}, 2: {"Gestor": None, "Ementa (Objetivo)": None}, 3: {"Unidade": None}})
    base = amostra()
    obrigatorios = [c for c, _ in projetos.CAMPOS_OBRIGATORIOS]
    esperado = int(base[obrigatorios].apply(lambda col: col.isna() | (col.astype(str).str.strip() == "")).any(axis=1).sum())
    assert esperado >= 3

    at = ir_para(entrar_codigo(PAULA[0]), "pendencias")
    assert not at.exception
    valores = {m.label: m.value for m in at.metric}
    assert valores["Projetos ativos"] == str(len(base))
    assert valores["Com campo obrigatório em branco"] == str(esperado)
    assert int(valores["Sem nenhuma pendência"]) <= len(base) - esperado


def test_pendencias_filtram_por_tipo_e_oferecem_exportacao(equipe):
    alterar_planilha({1: {"Município": None}, 2: {"Município": None}, 3: {"Unidade": None}})
    at = ir_para(entrar_codigo(PAULA[0]), "pendencias")
    assert len(at.get("download_button")) == 1
    at.multiselect(key="pend_categoria").set_value(["Município em branco"])
    at = at.run()
    legenda = " ".join(c.value for c in at.caption)
    esperado = int(amostra()["Município"].isna().sum())  # inclui o que já estava em branco no exemplo
    assert esperado >= 2 and f"{esperado} projeto(s) e {esperado} pendência(s)" in legenda
    at.text_input(key="pend_busca").set_value("texto-inexistente-xyz")
    at = at.run()
    assert any("Nenhuma pendência para o filtro" in c.value for c in at.caption)


def test_gestor_ve_aviso_com_o_que_falta_nos_proprios_projetos(equipe):
    base = amostra()
    n = int(base[base["Gestor"] == "Membro 2"].iloc[0]["Nº"])
    alterar_planilha({n: {"Município": None}})
    at = ir_para(entrar_codigo("membro.dois"), "meus")
    assert not at.exception
    aviso = next(w.value for w in at.warning if "campos obrigatórios em branco" in w.value)
    assert "município" in aviso and "Faltam:" in aviso


def test_quadro_e_pendencias_nao_tem_emoji_nem_a_palavra_auditoria(equipe):
    emoji = re.compile("[\U0001F300-\U0001FAFF☀-➿⬀-⯿️⃣]")
    at = entrar_codigo(PAULA[0])
    for pagina in ("quadro", "pendencias"):
        at = ir_para(at, pagina)
        assert not at.exception, pagina
        texto = textos_visiveis(at)
        assert not emoji.search(texto) and "audit" not in texto.lower(), pagina
